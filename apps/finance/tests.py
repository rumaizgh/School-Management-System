import io
import datetime
from decimal import Decimal

from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework import status
import openpyxl

from apps.academics.models import Institute, Batch, Fee, Payment
from apps.account.models import UserData
from apps.payroll.models import TeacherSalary, SalaryPayment
from apps.finance.models import InstitutionalTransaction


class FinanceModuleTests(TestCase):
    def setUp(self):
        self.client = APIClient()

        # Create Institute
        self.institute = Institute.objects.create(
            name="Greenwood High International",
            address="123 Education Blvd",
            established_date="2010-01-01"
        )

        # Create Admin User
        self.admin_user = UserData.objects.create(
            email="admin@greenwood.com",
            name="Admin Officer",
            user_type="admin",
            institute=self.institute
        )
        self.admin_user.set_password("Admin@12345")
        self.admin_user.save()

        # Create Teacher User
        self.teacher_user = UserData.objects.create(
            email="teacher@greenwood.com",
            name="Mr. Sharma",
            user_type="teacher",
            institute=self.institute
        )
        self.teacher_user.set_password("Teacher@12345")
        self.teacher_user.save()

        # Create Student User
        self.student_user = UserData.objects.create(
            email="student@greenwood.com",
            name="Rahul Roy",
            roll_no="STU-001",
            user_type="student",
            institute=self.institute
        )
        self.student_user.set_password("Student@12345")
        self.student_user.save()

        # Create Batch & Fee
        self.batch = Batch.objects.create(
            institute=self.institute,
            classs="10",
            year="2026-27"
        )
        self.fee = Fee.objects.create(
            institute=self.institute,
            student=self.student_user,
            batch=self.batch,
            amount=Decimal("50000.00"),
            description="Term 1 Tuition Fee"
        )
        # Create Fee Payment
        self.fee_payment = Payment.objects.create(
            fee=self.fee,
            amount=Decimal("30000.00"),
            payment_method="cash"
        )

        # Create Teacher Salary & Salary Payment
        now = timezone.now()
        month_str = now.strftime('%Y-%m')
        self.salary = TeacherSalary.objects.create(
            teacher=self.teacher_user,
            institute=self.institute,
            month=month_str,
            pay_type="monthly",
            base_salary=Decimal("40000.00")
        )
        self.salary_payment = SalaryPayment.objects.create(
            salary=self.salary,
            teacher=self.teacher_user,
            amount=Decimal("25000.00"),
            payment_method="bank_transfer",
            paid_on=now
        )

        # Authenticate client as admin by default
        self.client.force_authenticate(user=self.admin_user)

    # -------------------------------------------------------------
    # 1. Authentication & Permission Tests
    # -------------------------------------------------------------
    def test_unauthenticated_request_rejected(self):
        self.client.force_authenticate(user=None)
        res = self.client.get('/api/finance/balance-sheet/')
        self.assertIn(res.status_code, [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN])
        self.assertEqual(res.data.get('code'), 'PERMISSION_DENIED')

    def test_non_admin_request_rejected(self):
        self.client.force_authenticate(user=self.student_user)
        res = self.client.get('/api/finance/balance-sheet/')
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(res.data.get('code'), 'PERMISSION_DENIED')

    # -------------------------------------------------------------
    # 2. Transaction CRUD & Validation Tests
    # -------------------------------------------------------------
    def test_create_transaction_success(self):
        payload = {
            "type": "expense",
            "amount": 4500.00,
            "payment_method": "cash",
            "title": "Exam Pens, Pencils & Answer Sheets",
            "category": "stationery",
            "date": timezone.now().strftime("%Y-%m-%dT%H:%M:%SZ"),
            "payment_reference": None,
            "reference_number": "BILL-0982",
            "remarks": "Purchased 20 boxes of pens and answer booklets"
        }
        res = self.client.post('/api/finance/transactions/', data=payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data['status'], 'success')
        self.assertEqual(res.data['message'], 'Financial record added successfully.')
        self.assertTrue(res.data['data']['id'].startswith('tx-'))
        self.assertEqual(res.data['data']['amount'], 4500.00)
        self.assertEqual(res.data['data']['is_edited'], False)

    def test_create_transaction_validation_error_amount_zero_or_negative(self):
        payload = {
            "type": "expense",
            "amount": 0.00,
            "payment_method": "cash",
            "title": "Invalid Amount Item",
            "category": "stationery",
            "date": timezone.now().strftime("%Y-%m-%dT%H:%M:%SZ")
        }
        res = self.client.post('/api/finance/transactions/', data=payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data['status'], 'error')
        self.assertEqual(res.data['code'], 'VALIDATION_ERROR')
        self.assertEqual(res.data['message'], 'Amount must be greater than zero.')
        self.assertIn('amount', res.data['errors'])

    def test_create_transaction_validation_error_invalid_category(self):
        payload = {
            "type": "expense",
            "amount": 500.00,
            "payment_method": "cash",
            "title": "Invalid Category Item",
            "category": "unsupported_category_xyz",
            "date": timezone.now().strftime("%Y-%m-%dT%H:%M:%SZ")
        }
        res = self.client.post('/api/finance/transactions/', data=payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(res.data['code'], 'VALIDATION_ERROR')
        self.assertIn('category', res.data['errors'])

    def test_update_transaction_sets_is_edited_true(self):
        # First create record
        tx = InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Old Title",
            amount=Decimal("1000.00"),
            category="stationery",
            payment_method="cash",
            date=timezone.now(),
            created_by=self.admin_user,
            is_edited=False
        )

        update_payload = {
            "title": "Updated Title",
            "amount": 1200.00
        }
        # Update using tx-{id}
        res = self.client.patch(f'/api/finance/transactions/tx-{tx.id}/', data=update_payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['status'], 'success')
        self.assertEqual(res.data['message'], 'Financial record updated successfully.')
        self.assertEqual(res.data['data']['is_edited'], True)
        self.assertEqual(res.data['data']['amount'], 1200.00)
        self.assertEqual(res.data['data']['title'], 'Updated Title')

        tx.refresh_from_db()
        self.assertTrue(tx.is_edited)
        self.assertEqual(tx.amount, Decimal("1200.00"))

    def test_delete_transaction(self):
        tx = InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="To Delete",
            amount=Decimal("500.00"),
            category="stationery",
            payment_method="cash",
            date=timezone.now(),
            created_by=self.admin_user
        )

        res = self.client.delete(f'/api/finance/transactions/tx-{tx.id}/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['status'], 'success')
        self.assertEqual(res.data['message'], 'Financial record removed successfully.')
        self.assertFalse(InstitutionalTransaction.objects.filter(id=tx.id).exists())

    def test_record_not_found_404(self):
        res = self.client.get('/api/finance/transactions/tx-999/')
        # Detail view handles patch/delete/put
        res = self.client.delete('/api/finance/transactions/tx-999/')
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(res.data['code'], 'NOT_FOUND')
        self.assertEqual(res.data['message'], "Transaction with ID 'tx-999' does not exist.")

    # -------------------------------------------------------------
    # 3. Transaction Ledger Listing & Filtering Tests
    # -------------------------------------------------------------
    def test_transaction_list_and_filters(self):
        now = timezone.now()
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="income",
            title="Canteen Monthly Rent October",
            amount=Decimal("15000.00"),
            category="canteen_rent",
            payment_method="bank_transfer",
            reference_number="RENT-OCT-26",
            payment_reference="NEFT-782190",
            date=now,
            created_by=self.admin_user
        )
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Science Lab Chemicals",
            amount=Decimal("8000.00"),
            category="lab_supplies",
            payment_method="upi",
            reference_number="BILL-4421",
            payment_reference="UPI-992837",
            date=now,
            created_by=self.admin_user
        )

        # 1. Basic listing
        res = self.client.get('/api/finance/transactions/?period=all_time')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['status'], 'success')
        self.assertEqual(res.data['data']['total_records'], 2)
        self.assertEqual(len(res.data['data']['results']), 2)

        # 2. Filter by type=income
        res_income = self.client.get('/api/finance/transactions/?type=income&period=all_time')
        self.assertEqual(res_income.data['data']['total_records'], 1)
        self.assertEqual(res_income.data['data']['results'][0]['category'], 'canteen_rent')

        # 3. Filter by search=Chemicals
        res_search = self.client.get('/api/finance/transactions/?search=Chemicals&period=all_time')
        self.assertEqual(res_search.data['data']['total_records'], 1)
        self.assertEqual(res_search.data['data']['results'][0]['title'], 'Science Lab Chemicals')

        # 4. Filter by category=lab_supplies
        res_cat = self.client.get('/api/finance/transactions/?category=lab_supplies&period=all_time')
        self.assertEqual(res_cat.data['data']['total_records'], 1)

    # -------------------------------------------------------------
    # 4. Balance Sheet Calculations Tests (Zero Math)
    # -------------------------------------------------------------
    def test_balance_sheet_summary_calculations(self):
        now = timezone.now()
        # Add institutional income
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="income",
            title="Canteen Monthly Rent",
            amount=Decimal("20000.00"),
            category="canteen_rent",
            payment_method="bank_transfer",
            date=now,
            created_by=self.admin_user
        )
        # Add institutional expense
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Generator Fuel",
            amount=Decimal("5000.00"),
            category="maintenance",
            payment_method="cash",
            date=now,
            created_by=self.admin_user
        )

        # Expected:
        # fee_income = 30000.00
        # other_income = 20000.00
        # total_income = 50000.00
        # fee_percentage = (30000 / 50000) * 100 = 60.0%
        #
        # payroll_expense = 25000.00
        # other_expense = 5000.00
        # total_expense = 30000.00
        # payroll_percentage = (25000 / 30000) * 100 = 83.3%
        #
        # net_balance = 50000.00 - 30000.00 = 20000.00
        # balance_status = 'surplus'

        res = self.client.get('/api/finance/balance-sheet/?period=all_time')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['status'], 'success')
        d = res.data['data']

        self.assertEqual(d['total_income'], 50000.00)
        self.assertEqual(d['total_expense'], 30000.00)
        self.assertEqual(d['net_balance'], 20000.00)
        self.assertEqual(d['balance_status'], 'surplus')

        inc = d['breakdown']['income']
        self.assertEqual(inc['fee_income'], 30000.00)
        self.assertEqual(inc['other_income'], 20000.00)
        self.assertEqual(inc['fee_percentage'], 60.0)

        exp = d['breakdown']['expense']
        self.assertEqual(exp['payroll_expense'], 25000.00)
        self.assertEqual(exp['other_expense'], 5000.00)
        self.assertEqual(exp['payroll_percentage'], 83.3)

    # -------------------------------------------------------------
    # 5. Excel Statement Export Tests
    # -------------------------------------------------------------
    def test_export_statement_to_excel(self):
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Desk & Bench Repairs",
            amount=Decimal("3500.00"),
            category="maintenance",
            payment_method="cash",
            date=timezone.now(),
            created_by=self.admin_user
        )

        res = self.client.get('/api/finance/export/?period=all_time')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="Institution_Ledger_', res['Content-Disposition'])

        # Validate that the returned bytes are a valid readable Excel workbook with 2 sheets
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Balance Sheet Summary", wb.sheetnames)
        self.assertIn("Transactions Ledger", wb.sheetnames)

    # -------------------------------------------------------------
    # 6. Categories Metadata Tests
    # -------------------------------------------------------------
    def test_categories_and_payment_methods_metadata(self):
        res = self.client.get('/api/finance/categories/')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn('categories', res.data['data'])
        self.assertIn('payment_methods', res.data['data'])
        cats = [c['key'] for c in res.data['data']['categories']]
        self.assertIn('stationery', cats)
        self.assertIn('canteen_rent', cats)

    # -------------------------------------------------------------
    # 7. Multi-Tenant Isolation Tests
    # -------------------------------------------------------------
    def test_multitenant_isolation(self):
        other_institute = Institute.objects.create(
            name="Apex Academy",
            address="456 Secondary St",
            established_date="2015-01-01"
        )
        other_admin = UserData.objects.create(
            email="admin@apex.com",
            name="Apex Admin",
            user_type="admin",
            institute=other_institute
        )
        other_admin.set_password("ApexAdmin@123")
        other_admin.save()

        # Create Greenwood High transaction
        gw_tx = InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Greenwood Library Books",
            amount=Decimal("1500.00"),
            category="stationery",
            payment_method="cash",
            date=timezone.now(),
            created_by=self.admin_user
        )

        # Authenticate as Apex Admin
        self.client.force_authenticate(user=other_admin)

        # Apex should see 0 records in Greenwood
        res_list = self.client.get('/api/finance/transactions/?period=all_time')
        self.assertEqual(res_list.data['data']['total_records'], 0)

        # Apex cannot delete or patch Greenwood record
        res_del = self.client.delete(f'/api/finance/transactions/tx-{gw_tx.id}/')
        self.assertEqual(res_del.status_code, status.HTTP_404_NOT_FOUND)

        # Switch back to Greenwood Admin
        self.client.force_authenticate(user=self.admin_user)
        res_gw = self.client.get('/api/finance/transactions/?period=all_time')
        self.assertEqual(res_gw.data['data']['total_records'], 1)

    # -------------------------------------------------------------
    # 8. Custom Date Range Tests
    # -------------------------------------------------------------
    def test_custom_date_range_filtering(self):
        # Transaction in September 2026
        dt_sep = timezone.make_aware(datetime.datetime(2026, 9, 15, 12, 0, 0))
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="September Water Bill",
            amount=Decimal("1200.00"),
            category="utilities",
            payment_method="cash",
            date=dt_sep,
            created_by=self.admin_user
        )

        # Transaction in October 2026
        dt_oct = timezone.make_aware(datetime.datetime(2026, 10, 10, 12, 0, 0))
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="October Water Bill",
            amount=Decimal("1400.00"),
            category="utilities",
            payment_method="cash",
            date=dt_oct,
            created_by=self.admin_user
        )

        # Filter strictly for September 2026
        res = self.client.get('/api/finance/transactions/?start_date=2026-09-01&end_date=2026-09-30')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['data']['total_records'], 1)
        self.assertEqual(res.data['data']['results'][0]['title'], "September Water Bill")

    # -------------------------------------------------------------
    # 9. Daily (today / daily) Period Filtering Tests
    # -------------------------------------------------------------
    def test_daily_period_filtering(self):
        now = timezone.now()
        yesterday = now - datetime.timedelta(days=1)

        # Record today
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Today Snack Expense",
            amount=Decimal("350.00"),
            category="miscellaneous",
            payment_method="cash",
            date=now,
            created_by=self.admin_user
        )
        # Record yesterday
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="Yesterday Grocery",
            amount=Decimal("800.00"),
            category="miscellaneous",
            payment_method="cash",
            date=yesterday,
            created_by=self.admin_user
        )

        # 1. Balance sheet with period=today
        res_bs = self.client.get('/api/finance/balance-sheet/?period=today')
        self.assertEqual(res_bs.status_code, status.HTTP_200_OK)
        self.assertEqual(res_bs.data['data']['period_key'], 'today')
        self.assertEqual(res_bs.data['data']['breakdown']['expense']['other_expense'], 350.00)

        # 2. Balance sheet with alias period=daily
        res_daily = self.client.get('/api/finance/balance-sheet/?period=daily')
        self.assertEqual(res_daily.status_code, status.HTTP_200_OK)
        self.assertEqual(res_daily.data['data']['period_key'], 'today')

        # 3. Transaction list with period=today
        res_tx = self.client.get('/api/finance/transactions/?period=today')
        self.assertEqual(res_tx.status_code, status.HTTP_200_OK)
        self.assertEqual(res_tx.data['data']['total_records'], 1)
        self.assertEqual(res_tx.data['data']['results'][0]['title'], "Today Snack Expense")

        # 4. Excel export with period=today
        res_exp = self.client.get('/api/finance/export/?period=today')
        self.assertEqual(res_exp.status_code, status.HTTP_200_OK)
        self.assertIn('Institution_Ledger_Daily_', res_exp['Content-Disposition'])

    # -------------------------------------------------------------
    # 10. Weekly (this_week / weekly) Period Filtering Tests
    # -------------------------------------------------------------
    def test_weekly_period_filtering(self):
        now = timezone.now()
        monday = now.date() - datetime.timedelta(days=now.weekday())
        monday_dt = timezone.make_aware(datetime.datetime.combine(monday, datetime.time(10, 0)))
        two_weeks_ago = now - datetime.timedelta(days=14)

        # Record on Monday of current week
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="income",
            title="Monday Canteen Collection",
            amount=Decimal("5000.00"),
            category="canteen_rent",
            payment_method="cash",
            date=monday_dt,
            created_by=self.admin_user
        )
        # Record 2 weeks ago
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="income",
            title="Old Canteen Collection",
            amount=Decimal("5000.00"),
            category="canteen_rent",
            payment_method="cash",
            date=two_weeks_ago,
            created_by=self.admin_user
        )

        # 1. Balance sheet with period=this_week
        res_bs = self.client.get('/api/finance/balance-sheet/?period=this_week')
        self.assertEqual(res_bs.status_code, status.HTTP_200_OK)
        self.assertEqual(res_bs.data['data']['period_key'], 'this_week')
        self.assertEqual(res_bs.data['data']['breakdown']['income']['other_income'], 5000.00)

        # 2. Balance sheet with alias period=weekly
        res_weekly = self.client.get('/api/finance/balance-sheet/?period=weekly')
        self.assertEqual(res_weekly.status_code, status.HTTP_200_OK)
        self.assertEqual(res_weekly.data['data']['period_key'], 'this_week')

        # 3. Transactions list with period=this_week
        res_tx = self.client.get('/api/finance/transactions/?period=this_week')
        self.assertEqual(res_tx.status_code, status.HTTP_200_OK)
        self.assertEqual(res_tx.data['data']['total_records'], 1)
        self.assertEqual(res_tx.data['data']['results'][0]['title'], "Monday Canteen Collection")

        # 4. Excel export with period=this_week
        res_exp = self.client.get('/api/finance/export/?period=this_week')
        self.assertEqual(res_exp.status_code, status.HTTP_200_OK)
        self.assertIn('Institution_Ledger_Weekly_', res_exp['Content-Disposition'])

    # -------------------------------------------------------------
    # 11. Explicit period=custom Filtering Tests Across All 3 Endpoints
    # -------------------------------------------------------------
    def test_explicit_custom_period_across_all_endpoints(self):
        dt1 = timezone.make_aware(datetime.datetime(2026, 8, 10, 10, 0))
        InstitutionalTransaction.objects.create(
            institute=self.institute,
            type="expense",
            title="August Renovation",
            amount=Decimal("15000.00"),
            category="maintenance",
            payment_method="bank_transfer",
            date=dt1,
            created_by=self.admin_user
        )

        # 1. Balance sheet with period=custom
        res_bs = self.client.get('/api/finance/balance-sheet/?period=custom&start_date=2026-08-01&end_date=2026-08-31')
        self.assertEqual(res_bs.status_code, status.HTTP_200_OK)
        self.assertEqual(res_bs.data['data']['period_key'], 'custom')
        self.assertEqual(res_bs.data['data']['breakdown']['expense']['other_expense'], 15000.00)

        # 2. Transactions list with period=custom
        res_tx = self.client.get('/api/finance/transactions/?period=custom&start_date=2026-08-01&end_date=2026-08-31')
        self.assertEqual(res_tx.status_code, status.HTTP_200_OK)
        self.assertEqual(res_tx.data['data']['total_records'], 1)
        self.assertEqual(res_tx.data['data']['results'][0]['title'], "August Renovation")

        # 3. Excel export with period=custom
        res_exp = self.client.get('/api/finance/export/?period=custom&start_date=2026-08-01&end_date=2026-08-31')
        self.assertEqual(res_exp.status_code, status.HTTP_200_OK)
        self.assertIn('Institution_Ledger_Custom_2026-08-01_to_2026-08-31.xlsx', res_exp['Content-Disposition'])

