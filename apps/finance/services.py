import io
import datetime
from calendar import monthrange
from decimal import Decimal

from django.db.models import Sum, Q
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.academics.models import Payment, Payroll
from apps.payroll.models import SalaryPayment
from .models import InstitutionalTransaction


CATEGORIES_METADATA = [
    {
        "key": "stationery",
        "type": "expense",
        "label": "Stationery",
        "examples": "Pens, paper bundles, answer booklets, printer cartridges"
    },
    {
        "key": "utilities",
        "type": "expense",
        "label": "Utilities & Internet",
        "examples": "Broadband fiber bills, water, telephone"
    },
    {
        "key": "maintenance",
        "type": "expense",
        "label": "Maintenance & Repairs",
        "examples": "Generator diesel, water cooler repair, electrical service"
    },
    {
        "key": "lab_supplies",
        "type": "expense",
        "label": "Lab Supplies",
        "examples": "Beakers, chemicals, physics apparatus"
    },
    {
        "key": "event",
        "type": "expense",
        "label": "Events & Functions",
        "examples": "Annual day stage, sound system, sports day tents"
    },
    {
        "key": "canteen_rent",
        "type": "income",
        "label": "Canteen Rent",
        "examples": "Monthly rent paid by cafeteria contractor"
    },
    {
        "key": "donation",
        "type": "income",
        "label": "Donations & Grants",
        "examples": "Alumni donations, CSR educational grants"
    },
    {
        "key": "fine",
        "type": "income",
        "label": "Fines & Penalties",
        "examples": "Library book delay fines, late fees"
    },
    {
        "key": "miscellaneous",
        "type": "both",
        "label": "Miscellaneous",
        "examples": "General or unclassified records"
    },
]

PAYMENT_METHODS_METADATA = [
    {"key": "cash", "label": "Cash"},
    {"key": "upi", "label": "UPI"},
    {"key": "bank_transfer", "label": "Bank Transfer (NEFT / RTGS / IMPS)"},
    {"key": "check", "label": "Cheque"},
]


def resolve_period_dates(period='this_month', start_date_str=None, end_date_str=None):
    """
    Computes exact timezone-aware start and end timestamps and labels
    for the requested period filter.
    Supported periods:
    - today / daily: current day from 00:00:00 to 23:59:59
    - this_week / weekly: Monday of current week up to today (23:59:59)
    - this_month: 1st of month to end of month
    - this_term: current term
    - this_academic_year: current academic year
    - all_time: all records
    - custom: between start_date 00:00:00 and end_date 23:59:59 (YYYY-MM-DD)
    """
    now = timezone.now()
    period_key = (period or 'this_month').lower().strip()

    # 1. Custom date range (if explicitly requested as 'custom' or if dates are provided)
    if period_key == 'custom' or start_date_str or end_date_str:
        start_dt = None
        end_dt = None
        if start_date_str:
            try:
                d = datetime.date.fromisoformat(str(start_date_str).strip())
                start_dt = timezone.make_aware(datetime.datetime.combine(d, datetime.time.min))
            except (ValueError, TypeError):
                pass
        if end_date_str:
            try:
                d = datetime.date.fromisoformat(str(end_date_str).strip())
                end_dt = timezone.make_aware(datetime.datetime.combine(d, datetime.time.max))
            except (ValueError, TypeError):
                pass
        return {
            "period": f"{start_date_str or 'Start'} to {end_date_str or 'End'}",
            "period_key": "custom",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 2. Daily: today / daily (Current day from 00:00:00 to 23:59:59)
    if period_key in ['today', 'daily']:
        today_date = now.date()
        start_dt = timezone.make_aware(datetime.datetime.combine(today_date, datetime.time.min))
        end_dt = timezone.make_aware(datetime.datetime.combine(today_date, datetime.time.max))
        return {
            "period": "Today",
            "period_key": "today",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 3. Weekly: this_week / weekly (Monday of current week up to today)
    elif period_key in ['this_week', 'weekly']:
        monday_date = now.date() - datetime.timedelta(days=now.weekday())
        start_dt = timezone.make_aware(datetime.datetime.combine(monday_date, datetime.time.min))
        end_dt = timezone.make_aware(datetime.datetime.combine(now.date(), datetime.time.max))
        return {
            "period": "This Week",
            "period_key": "this_week",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 4. Monthly: this_month
    elif period_key == 'this_month':
        year, month = now.year, now.month
        last_day = monthrange(year, month)[1]
        start_dt = timezone.make_aware(datetime.datetime(year, month, 1, 0, 0, 0))
        end_dt = timezone.make_aware(datetime.datetime(year, month, last_day, 23, 59, 59))
        return {
            "period": "This Month",
            "period_key": "this_month",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 5. Term: this_term
    elif period_key == 'this_term':
        year, month = now.year, now.month
        if 6 <= month <= 10:
            start_dt = timezone.make_aware(datetime.datetime(year, 6, 1, 0, 0, 0))
            end_dt = timezone.make_aware(datetime.datetime(year, 10, 31, 23, 59, 59))
        elif month in [11, 12]:
            start_dt = timezone.make_aware(datetime.datetime(year, 11, 1, 0, 0, 0))
            end_dt = timezone.make_aware(datetime.datetime(year + 1, 3, 31, 23, 59, 59))
        elif 1 <= month <= 3:
            start_dt = timezone.make_aware(datetime.datetime(year - 1, 11, 1, 0, 0, 0))
            end_dt = timezone.make_aware(datetime.datetime(year, 3, 31, 23, 59, 59))
        else:
            start_dt = timezone.make_aware(datetime.datetime(year, 4, 1, 0, 0, 0))
            end_dt = timezone.make_aware(datetime.datetime(year, 5, 31, 23, 59, 59))
        return {
            "period": "This Term",
            "period_key": "this_term",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 6. Academic Year: this_academic_year
    elif period_key == 'this_academic_year':
        year, month = now.year, now.month
        if month >= 6:
            start_year = year
            end_year = year + 1
        else:
            start_year = year - 1
            end_year = year
        start_dt = timezone.make_aware(datetime.datetime(start_year, 6, 1, 0, 0, 0))
        end_dt = timezone.make_aware(datetime.datetime(end_year, 5, 31, 23, 59, 59))
        return {
            "period": "This Academic Year",
            "period_key": "this_academic_year",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }

    # 7. All Time: all_time
    elif period_key == 'all_time':
        return {
            "period": "All Time",
            "period_key": "all_time",
            "start_datetime": None,
            "end_datetime": None,
        }

    else:
        # Default fallback to this_month
        year, month = now.year, now.month
        last_day = monthrange(year, month)[1]
        start_dt = timezone.make_aware(datetime.datetime(year, month, 1, 0, 0, 0))
        end_dt = timezone.make_aware(datetime.datetime(year, month, last_day, 23, 59, 59))
        return {
            "period": "This Month",
            "period_key": "this_month",
            "start_datetime": start_dt,
            "end_datetime": end_dt,
        }


def get_balance_sheet_summary(institute, period_info):
    """
    Computes zero-frontend-math aggregated figures for the institutional balance sheet:
    - total_income = fee_income (Student Payments) + other_income (Institutional Transactions 'income')
    - total_expense = payroll_expense (Salary Payments) + other_expense (Institutional Transactions 'expense')
    - net_balance = total_income - total_expense
    - balance_status = 'surplus' if net_balance >= 0 else 'deficit'
    """
    start_dt = period_info.get('start_datetime')
    end_dt = period_info.get('end_datetime')

    # 1. Student Fee Income (academics.Payment)
    fee_qs = Payment.objects.filter(
        Q(fee__institute=institute) | Q(fee__student__institute=institute)
    ) if institute else Payment.objects.all()

    if start_dt:
        fee_qs = fee_qs.filter(paid_on__gte=start_dt.date())
    if end_dt:
        fee_qs = fee_qs.filter(paid_on__lte=end_dt.date())

    fee_income_dec = fee_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # 2. Other Institutional Income
    other_income_qs = InstitutionalTransaction.objects.filter(type='income')
    if institute:
        other_income_qs = other_income_qs.filter(institute=institute)
    if start_dt:
        other_income_qs = other_income_qs.filter(date__gte=start_dt)
    if end_dt:
        other_income_qs = other_income_qs.filter(date__lte=end_dt)

    other_income_dec = other_income_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # Total Income & Breakdown
    total_income_dec = fee_income_dec + other_income_dec
    if total_income_dec > Decimal('0.00'):
        fee_percentage = round(float(fee_income_dec) / float(total_income_dec) * 100, 1)
    else:
        fee_percentage = 0.0

    # 3. Staff & Teacher Payroll Expense
    # From apps.payroll.SalaryPayment
    payroll_qs = SalaryPayment.objects.filter(
        Q(salary__institute=institute) | Q(teacher__institute=institute)
    ) if institute else SalaryPayment.objects.all()

    if start_dt:
        payroll_qs = payroll_qs.filter(paid_on__gte=start_dt)
    if end_dt:
        payroll_qs = payroll_qs.filter(paid_on__lte=end_dt)

    salary_payroll_dec = payroll_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # Legacy academics.Payroll (if any exists)
    legacy_payroll_qs = Payroll.objects.filter(institute=institute) if institute else Payroll.objects.all()
    if start_dt:
        legacy_payroll_qs = legacy_payroll_qs.filter(created_at__gte=start_dt.date())
    if end_dt:
        legacy_payroll_qs = legacy_payroll_qs.filter(created_at__lte=end_dt.date())
    legacy_payroll_dec = legacy_payroll_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    payroll_expense_dec = salary_payroll_dec + legacy_payroll_dec

    # 4. Other Institutional Expenses
    other_expense_qs = InstitutionalTransaction.objects.filter(type='expense')
    if institute:
        other_expense_qs = other_expense_qs.filter(institute=institute)
    if start_dt:
        other_expense_qs = other_expense_qs.filter(date__gte=start_dt)
    if end_dt:
        other_expense_qs = other_expense_qs.filter(date__lte=end_dt)

    other_expense_dec = other_expense_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')

    # Total Expense & Breakdown
    total_expense_dec = payroll_expense_dec + other_expense_dec
    if total_expense_dec > Decimal('0.00'):
        payroll_percentage = round(float(payroll_expense_dec) / float(total_expense_dec) * 100, 1)
    else:
        payroll_percentage = 0.0

    # 5. Net Balance & Status
    net_balance_dec = total_income_dec - total_expense_dec
    balance_status = 'surplus' if net_balance_dec >= Decimal('0.00') else 'deficit'

    return {
        "period": period_info['period'],
        "period_key": period_info['period_key'],
        "start_date": start_dt.strftime('%Y-%m-%dT%H:%M:%SZ') if start_dt else None,
        "end_date": end_dt.strftime('%Y-%m-%dT%H:%M:%SZ') if end_dt else None,
        "total_income": float(total_income_dec),
        "total_expense": float(total_expense_dec),
        "net_balance": float(net_balance_dec),
        "balance_status": balance_status,
        "breakdown": {
            "income": {
                "fee_income": float(fee_income_dec),
                "other_income": float(other_income_dec),
                "fee_percentage": float(fee_percentage),
            },
            "expense": {
                "payroll_expense": float(payroll_expense_dec),
                "other_expense": float(other_expense_dec),
                "payroll_percentage": float(payroll_percentage),
            }
        }
    }


def generate_financial_excel_report(institute, period_info, type_filter, transactions_qs):
    """
    Generates an official multi-sheet Excel (.xlsx) financial statement workbook:
    - Sheet 1: Executive Balance Sheet Summary
    - Sheet 2: Itemized Transactions Ledger
    """
    summary = get_balance_sheet_summary(institute, period_info)

    wb = openpyxl.Workbook()

    # Common styles
    font_family = "Segoe UI"
    navy_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    blue_sub_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    emerald_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    rose_fill = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    zebra_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    title_font = Font(name=font_family, size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name=font_family, size=11, bold=False, color="475569")
    meta_font = Font(name=font_family, size=10, italic=True, color="64748B")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=11, bold=True, color="0F172A")
    regular_font = Font(name=font_family, size=10, color="1E293B")
    surplus_font = Font(name=font_family, size=12, bold=True, color="047857")
    deficit_font = Font(name=font_family, size=12, bold=True, color="B91C1C")

    thin_border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )
    header_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=Side(style='medium', color='1E3A8A')
    )

    currency_fmt = "₹#,##0.00"
    pct_fmt = "0.0%"

    # -------------------------------------------------------------
    # SHEET 1: Balance Sheet Summary
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Balance Sheet Summary"
    ws1.views.sheetView[0].showGridLines = True

    # Title Block
    inst_name = institute.name if institute else "School Management System"
    ws1["A1"] = inst_name.upper()
    ws1["A1"].font = title_font

    ws1["A2"] = "INSTITUTIONAL BALANCE SHEET & FINANCIAL STATEMENT"
    ws1["A2"].font = subtitle_font

    now_str = timezone.now().strftime("%d %b %Y, %I:%M %p UTC")
    ws1["A3"] = f"Period: {summary['period']}  |  Generated On: {now_str}  |  Currency: INR (₹)"
    ws1["A3"].font = meta_font

    # Section 1: Revenue & Income
    ws1["A5"] = "1. REVENUE & INCOME STREAMS"
    ws1["A5"].font = bold_font

    income_headers = ["Revenue Stream", "Source Type", "Amount (INR)", "Share of Total"]
    for col_idx, h in enumerate(income_headers, start=1):
        cell = ws1.cell(row=6, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [3, 4] else "left", vertical="center")
        cell.border = header_border
    ws1.row_dimensions[6].height = 24

    fee_income = summary['breakdown']['income']['fee_income']
    other_income = summary['breakdown']['income']['other_income']
    total_income = summary['total_income']
    fee_pct = summary['breakdown']['income']['fee_percentage'] / 100.0
    other_pct = (100.0 - summary['breakdown']['income']['fee_percentage']) / 100.0 if total_income > 0 else 0.0

    income_rows = [
        ("Student Fee Collection", "Student Payment Module", fee_income, fee_pct),
        ("Other Institutional Income", "Institutional Ledger (Rent, Grants, Fines)", other_income, other_pct),
    ]

    for r_idx, row_data in enumerate(income_rows, start=7):
        ws1.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if c_idx == 3:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 4:
                cell.number_format = pct_fmt
                cell.alignment = Alignment(horizontal="right")

    # Total Income Row
    ws1.row_dimensions[9].height = 22
    t_inc_cells = [
        ws1.cell(row=9, column=1, value="TOTAL REVENUE / INCOME"),
        ws1.cell(row=9, column=2, value="Sum of All Income"),
        ws1.cell(row=9, column=3, value=total_income),
        ws1.cell(row=9, column=4, value=1.0 if total_income > 0 else 0.0),
    ]
    for c_idx, c in enumerate(t_inc_cells, start=1):
        c.font = bold_font
        c.fill = emerald_fill
        c.border = cell_border
        if c_idx == 3:
            c.number_format = currency_fmt
            c.alignment = Alignment(horizontal="right")
        elif c_idx == 4:
            c.number_format = pct_fmt
            c.alignment = Alignment(horizontal="right")

    # Section 2: Expenditure & Expenses
    ws1["A11"] = "2. EXPENDITURE & EXPENSE STREAMS"
    ws1["A11"].font = bold_font

    expense_headers = ["Expense Category", "Source Type", "Amount (INR)", "Share of Total"]
    for col_idx, h in enumerate(expense_headers, start=1):
        cell = ws1.cell(row=12, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [3, 4] else "left", vertical="center")
        cell.border = header_border
    ws1.row_dimensions[12].height = 24

    payroll_expense = summary['breakdown']['expense']['payroll_expense']
    other_expense = summary['breakdown']['expense']['other_expense']
    total_expense = summary['total_expense']
    payroll_pct = summary['breakdown']['expense']['payroll_percentage'] / 100.0
    other_exp_pct = (100.0 - summary['breakdown']['expense']['payroll_percentage']) / 100.0 if total_expense > 0 else 0.0

    expense_rows = [
        ("Teacher & Staff Payroll", "Payroll Disbursal Module", payroll_expense, payroll_pct),
        ("Other Institutional Expenses", "Institutional Ledger (Supplies, Bills, Repairs)", other_expense, other_exp_pct),
    ]

    for r_idx, row_data in enumerate(expense_rows, start=13):
        ws1.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if c_idx == 3:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c_idx == 4:
                cell.number_format = pct_fmt
                cell.alignment = Alignment(horizontal="right")

    # Total Expense Row
    ws1.row_dimensions[15].height = 22
    t_exp_cells = [
        ws1.cell(row=15, column=1, value="TOTAL EXPENDITURE / EXPENSE"),
        ws1.cell(row=15, column=2, value="Sum of All Expenses"),
        ws1.cell(row=15, column=3, value=total_expense),
        ws1.cell(row=15, column=4, value=1.0 if total_expense > 0 else 0.0),
    ]
    for c_idx, c in enumerate(t_exp_cells, start=1):
        c.font = bold_font
        c.fill = rose_fill
        c.border = cell_border
        if c_idx == 3:
            c.number_format = currency_fmt
            c.alignment = Alignment(horizontal="right")
        elif c_idx == 4:
            c.number_format = pct_fmt
            c.alignment = Alignment(horizontal="right")

    # Section 3: Net Balance Position
    ws1["A17"] = "3. EXECUTIVE NET POSITION (PROFIT IN HAND)"
    ws1["A17"].font = bold_font

    pos_headers = ["Metric", "Description", "Value (INR)", "Status Indicator"]
    for col_idx, h in enumerate(pos_headers, start=1):
        cell = ws1.cell(row=18, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [3, 4] else "left", vertical="center")
        cell.border = header_border
    ws1.row_dimensions[18].height = 24

    net_balance = summary['net_balance']
    status_label = summary['balance_status'].upper()

    ws1.row_dimensions[19].height = 24
    c_p1 = ws1.cell(row=19, column=1, value="NET BALANCE")
    c_p2 = ws1.cell(row=19, column=2, value="Total Income Minus Total Expenses")
    c_p3 = ws1.cell(row=19, column=3, value=net_balance)
    c_p4 = ws1.cell(row=19, column=4, value=f"● {status_label}")

    c_p1.font = bold_font
    c_p2.font = regular_font
    c_p3.font = surplus_font if net_balance >= 0 else deficit_font
    c_p4.font = surplus_font if net_balance >= 0 else deficit_font

    c_p3.number_format = currency_fmt
    c_p3.alignment = Alignment(horizontal="right")
    c_p4.alignment = Alignment(horizontal="center")

    for c in [c_p1, c_p2, c_p3, c_p4]:
        c.fill = emerald_fill if net_balance >= 0 else rose_fill
        c.border = cell_border

    # Auto-adjust column widths for Sheet 1
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 18)

    # -------------------------------------------------------------
    # SHEET 2: Itemized Transactions Ledger
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Transactions Ledger")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"{inst_name.upper()} — TRANSACTIONS LEDGER"
    ws2["A1"].font = title_font

    type_disp = type_filter.upper() if type_filter != 'all' else 'ALL (Income & Expense)'
    ws2["A2"] = f"Filter Scope: {type_disp}  |  Period: {summary['period']}"
    ws2["A2"].font = subtitle_font

    ledger_headers = [
        "Tx ID",
        "Date (UTC)",
        "Type",
        "Particulars / Title",
        "Category",
        "Payment Mode",
        "Reference No",
        "Payment Reference",
        "Amount (INR)",
        "Edited",
        "Remarks",
    ]

    for col_idx, h in enumerate(ledger_headers, start=1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 3, 6, 10] else "left", vertical="center")
        cell.border = header_border
    ws2.row_dimensions[4].height = 24

    for row_num, tx in enumerate(transactions_qs, start=5):
        ws2.row_dimensions[row_num].height = 19
        bg_fill = zebra_alt_fill if row_num % 2 == 0 else PatternFill(fill_type=None)

        type_display = tx.type.upper()
        type_color = "047857" if tx.type == 'income' else "B91C1C"

        row_values = [
            f"tx-{tx.id}",
            tx.date.strftime("%Y-%m-%d %H:%M") if tx.date else "",
            type_display,
            tx.title,
            tx.category.replace('_', ' ').title(),
            tx.payment_method.replace('_', ' ').title(),
            tx.reference_number or "-",
            tx.payment_reference or "-",
            float(tx.amount),
            "Yes" if tx.is_edited else "No",
            tx.remarks or "-",
        ]

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws2.cell(row=row_num, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if bg_fill.fill_type:
                cell.fill = bg_fill

            if col_idx in [1, 2, 6, 10]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                cell.font = Font(name=font_family, size=10, bold=True, color=type_color)
            elif col_idx == 9:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
                cell.font = Font(name=font_family, size=10, bold=True, color=type_color)

    # Auto-adjust column widths for Sheet 2
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
